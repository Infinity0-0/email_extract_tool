# -*- coding: utf-8 -*-
"""
EMAIL EXTRACTOR  (PDF links -> Excel with emails)

Kya karta hai:
  1. Neeche LINKS me jitne link daloge, sabko ek-ek karke kholta hai
  2. Har link se PDF dhundh kar download karta hai
  3. PDF ke har page ko padhta hai aur emails nikalta hai
  4. Sab kuch ek Excel file me save karta hai

Chalane ke liye:   python email_extractor.py
"""

import hashlib
import os
import re
import sys
import time
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import unquote, urljoin, urlparse

import pandas as pd
import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from tqdm import tqdm
from urllib3.util.retry import Retry

try:
    import fitz  # PyMuPDF
except ImportError:
    sys.exit("PyMuPDF install nahi hai. Terminal me chalao:  pip install pymupdf")

# =============================================================================
#  STEP 1:  YAHAN APNE LINKS PASTE KARO
#  Format: har link "quotes" me aur aakhir me comma (,)
#  Aur links add karne ho to bas niche naye line me isi format me jod do.
# =============================================================================
LINKS = []

# =============================================================================
#  STEP 2 (optional):  SETTINGS
# =============================================================================
OUTPUT_FILE = "emails_output.xlsx"   # result Excel ka naam
PDF_FOLDER = "downloaded_pdfs"       # PDFs yaha save hongi
WORKERS = 3                          # ek saath kitne download (zyada karoge to site block kar sakti hai)
TIMEOUT = 60                         # ek link ke liye max seconds
MAX_PAGES = 0                        # 0 = poori PDF padho, 3 = sirf pehle 3 pages
DELETE_PDF_AFTER = False             # True = email nikalne ke baad PDF delete (disk bachti hai)
SAB_KUCH_DO = True                   # True = tuta-futa/adhoora email bhi list me aaye (kuch bhi hata nahi)
                                     # False = adhoore emails khud hat jayein


# =============================================================================
#  NEECHE KUCH BADALNE KI ZARURAT NAHI
# =============================================================================

# ---------- Email nikalne ka logic ----------
RANK = {"High": 0, "Medium": 1}
GENERIC_TLDS = {
    "com", "org", "net", "edu", "gov", "mil", "int", "info", "biz", "name", "pro",
    "tech", "online", "site", "app", "dev", "xyz", "cloud", "science", "academy",
    "university", "institute", "center", "centre", "school", "email", "digital",
    "global", "systems", "press", "media", "lab", "labs", "health", "research",
}
SECOND_LEVEL = {"ac", "co", "edu", "gov", "org", "com", "net", "res", "nic", "sch"}
COMMON_CC = {
    "in", "uk", "au", "nz", "za", "sg", "my", "pk", "bd", "lk", "np", "jp", "cn", "kr",
    "hk", "tw", "br", "mx", "ar", "ae", "sa", "tr", "ng", "ke", "ph", "id", "th", "vn",
    "de", "fr", "it", "es", "nl", "se", "no", "dk", "fi", "pl", "ru", "ca", "us", "ie", "ch",
}
BAD_TLDS = {
    "png", "jpg", "jpeg", "gif", "svg", "webp", "pdf", "css", "js", "html", "htm",
    "doc", "docx", "xls", "xlsx", "zip", "txt", "csv", "json", "xml",
}
PLAIN_RE = re.compile(
    r"(?<![A-Za-z0-9._%+\-])"
    r"([A-Za-z0-9][A-Za-z0-9._%+\-]*)@"
    r"((?:[A-Za-z0-9](?:[A-Za-z0-9\-]*[A-Za-z0-9])?\.)+[A-Za-z]{2,24})"
)
BRACE_RE = re.compile(
    r"[\{\[\(]\s*([A-Za-z0-9._%+\-]+(?:\s*[,;]\s*[A-Za-z0-9._%+\-]+)*)\s*[\}\]\)]"
    r"\s*@\s*((?:[A-Za-z0-9](?:[A-Za-z0-9\-]*[A-Za-z0-9])?\.)+[A-Za-z]{2,24})"
)
OBF_AT = re.compile(r"\s*[\[\(\{<]\s*(?:at|AT|At)\s*[\]\)\}>]\s*|\s+AT\s+")
OBF_DOT = re.compile(r"\s*[\[\(\{<]\s*(?:dot|DOT|Dot)\s*[\]\)\}>]\s*|\s+DOT\s+")
ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(s):
    return ILLEGAL_XML.sub("", s) if isinstance(s, str) else s


def fix_domain(domain):
    """'univ.eduAbstract' -> 'univ.edu' (PDF me chipke shabd hatata hai)."""
    labels = domain.strip(".-").lower().split(".")
    while len(labels) >= 2:
        tld = labels[-1]
        if tld in BAD_TLDS:
            return None
        ok = tld in GENERIC_TLDS or len(tld) == 2
        if not ok:
            for g in sorted(GENERIC_TLDS, key=len, reverse=True):
                if tld.startswith(g):
                    labels[-1], ok = g, True
                    break
        if not ok and len(labels) >= 3 and labels[-2] in SECOND_LEVEL and tld[:2] in COMMON_CC:
            labels[-1], ok = tld[:2], True
        if ok:
            if any((not l) or l.startswith("-") or l.endswith("-") for l in labels):
                return None
            d = ".".join(labels)
            return d if len(d) <= 253 else None
        labels.pop()
    return None


def looks_cut(domain, sep):
    if sep == "-":
        return True
    tld = domain.strip(".").split(".")[-1].lower()
    return not (tld in GENERIC_TLDS or (len(tld) == 2 and tld not in SECOND_LEVEL))


def clean_local(local):
    local = local.strip("._%+-")
    if not (1 <= len(local) <= 64) or ".." in local or not re.search(r"[A-Za-z0-9]", local):
        return None
    return local


def build_email(local, domain):
    local, domain = clean_local(local), fix_domain(domain)
    return f"{local}@{domain}".lower() if local and domain else None


def snippet(text, start, end):
    return clean(re.sub(r"\s+", " ", text[max(0, start - 70): end + 30]).strip())


def repair_wraps(text):
    """Line ke end me toote emails jodta hai: 'john@univer-' + 'sity.edu'."""
    lines = text.split("\n")
    out = []
    for i in range(len(lines)):
        cur = lines[i].rstrip()
        if i + 1 < len(lines) and lines[i + 1].strip():
            nxt = lines[i + 1].strip()
            tok = nxt.split()[0]
            rest = nxt[len(tok):].strip()
            last = cur.split()[-1] if cur.split() else ""
            mb = re.search(r"@([A-Za-z0-9.\-]*)([.\-])$", cur)
            domain_break = bool(mb) and looks_cut(mb.group(1), mb.group(2))
            local_break = "@" in tok and re.fullmatch(r"[A-Za-z0-9._%+\-]*[A-Za-z0-9][._\-]", last or "x")
            if domain_break or local_break:
                cur = (cur[:-1] if cur.endswith("-") else cur) + tok
                lines[i + 1] = rest
        out.append(cur)
    return "\n".join(out)


def extract_from_text(text):
    """Text se emails -> {email: {conf, method, context}}"""
    found = {}

    def add(email, conf, method, ctx):
        old = found.get(email)
        if old is None or RANK[conf] < RANK[old["conf"]]:
            found[email] = {"conf": conf, "method": method, "context": ctx}

    def scan(t, conf, method, check_cut=False):
        for m in PLAIN_RE.finditer(t):
            e = build_email(m.group(1), m.group(2))
            if not e:
                continue
            c, meth = conf, method
            if check_cut:
                tail = re.match(r"([.\-])\s*\n\s*[a-z]{2,24}\b", t[m.end(): m.end() + 40])
                head = t[max(0, m.start() - 6): m.start()]
                cut_after = bool(tail) and looks_cut(m.group(2), tail.group(1))
                cut_before = bool(re.search(r"[A-Za-z0-9][\-_]\s*\n\s*$", head))
                if cut_after or cut_before:
                    c, meth = "Medium", "text (line-break pe toota ho sakta hai - check karo)"
            add(e, c, meth, snippet(t, m.start(), m.end()))

    scan(text, "High", "text", check_cut=True)

    for m in BRACE_RE.finditer(text):               # {a, b}@univ.edu
        dom = fix_domain(m.group(2))
        if not dom:
            continue
        for name in re.split(r"[,;\s]+", m.group(1)):
            e = build_email(name, dom)
            if e:
                add(e, "High", "brace-expanded", snippet(text, m.start(), m.end()))

    scan(repair_wraps(text), "Medium", "line-wrap repaired")
    scan(re.sub(r"[ \t]*@[ \t]*", "@", text), "Medium", "spaced @ repaired")
    obf = OBF_DOT.sub(".", OBF_AT.sub("@", text))   # name [at] domain [dot] com
    if obf != text:
        scan(obf, "Medium", "obfuscated (at/dot)")
    return found


def drop_broken(rows):
    """Line-break se toote hue adhoore emails hata do, agar poora email already mil gaya ho.
    (jaise 'mar@abc.org' hata do jab 'raj.kumar@abc.org' mil chuka ho)"""
    emails = [r["email"] for r in rows]
    keep = []
    for r in rows:
        if r["conf"] == "Medium" and "check karo" in r["method"]:
            l, d = r["email"].split("@")
            broken = False
            for f in emails:
                if f == r["email"]:
                    continue
                fl, fd = f.split("@")
                if (fd == d and fl.endswith(l) and len(fl) > len(l)) or \
                   (fl == l and fd.startswith(d) and len(fd) > len(d)):
                    broken = True
                    break
            if broken:
                continue
        keep.append(r)
    return keep


# ---------- PDF padhna ----------
def process_pdf(path):
    res = {"status": "OK", "note": "", "total_pages": 0, "pages_scanned": 0, "rows": []}
    try:
        doc = fitz.open(path)
        if doc.needs_pass and not doc.authenticate(""):
            raise ValueError("PDF password-protected hai")
        total = doc.page_count
        limit = min(total, MAX_PAGES) if MAX_PAGES else total
        pages = []
        for i in range(limit):
            page = doc[i]
            text = page.get_text("text") or ""
            links = []
            try:
                for l in page.get_links():
                    uri = l.get("uri") or ""
                    if uri.lower().startswith("mailto:"):
                        try:
                            ltxt = re.sub(r"\s+", " ", page.get_textbox(l["from"])).strip()
                        except Exception:
                            ltxt = ""
                        links.append((unquote(uri[7:].split("?")[0]), ltxt))
            except Exception:
                pass
            pages.append((i + 1, text, links))
        doc.close()
    except Exception as e:
        res.update(status="Failed", note=f"PDF read error: {e}")
        return res

    res["total_pages"], res["pages_scanned"] = total, len(pages)
    agg = OrderedDict()

    def merge(email, conf, method, page_no, ctx, text=""):
        low = text.lower()
        pos = low.find(email)
        if pos < 0:
            pos = low.find(email.split("@")[0])
        key = (page_no, pos if pos >= 0 else 10 ** 9)
        r = agg.get(email)
        if r is None:
            agg[email] = {"email": email, "conf": conf, "method": method, "order": key,
                          "pages": {page_no}, "count": 1, "context": ctx}
        else:
            r["pages"].add(page_no)
            r["count"] += 1
            r["order"] = min(r["order"], key)
            if RANK[conf] < RANK[r["conf"]]:
                r.update(conf=conf, method=method, context=ctx)

    chars = 0
    for page_no, text, links in pages:
        chars += len(text.strip())
        for blob, ltxt in links:                     # mailto: links (sabse reliable)
            for part in re.split(r"[,;]", blob):
                m = PLAIN_RE.search(part.strip())
                e = build_email(m.group(1), m.group(2)) if m else None
                if e:
                    merge(e, "High", "mailto-link", page_no, clean(ltxt), text)
        for email, info in extract_from_text(text).items():
            merge(email, info["conf"], info["method"], page_no, info["context"], text)

    if chars < 50 and not agg:
        res["status"], res["note"] = "No text", "Shayad scanned (image) PDF hai - text nahi mila"
    rows = list(agg.values())
    if not SAB_KUCH_DO:
        rows = drop_broken(rows)
    rows.sort(key=lambda r: r["order"])
    res["rows"] = [{**r, "pages": ", ".join(map(str, sorted(r["pages"])))} for r in rows]
    return res


# ---------- Link se PDF dhundhna + download ----------
def make_session():
    s = requests.Session()
    retry = Retry(total=3, backoff_factor=1.5, status_forcelist=[429, 500, 502, 503, 504],
                  allowed_methods=["GET"])
    s.mount("http://", HTTPAdapter(max_retries=retry))
    s.mount("https://", HTTPAdapter(max_retries=retry))
    s.headers.update({
        "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                       "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"),
        "Accept": "*/*",
    })
    return s


def discover(session, link):
    """Ek link -> [(pdf_url, article_title)]
    - seedha .pdf link ho to wahi
    - journal (OJS) ka 'article/view' page ho to uska Download PDF link"""
    if urlparse(link).path.lower().endswith(".pdf"):
        return [(link, "")]
    r = session.get(link, timeout=TIMEOUT)
    r.raise_for_status()
    if b"%PDF" in r.content[:1024] and "html" not in r.headers.get("Content-Type", "").lower():
        return [(link, "")]
    soup = BeautifulSoup(r.content, "html.parser")
    title = re.sub(r"^View of\s+", "", soup.title.get_text(strip=True)) if soup.title else ""
    out, seen = [], set()
    for tag in soup.find_all(["a", "iframe", "embed", "object"]):
        href = tag.get("href") or tag.get("src") or tag.get("data") or ""
        h = href.lower()
        if ".pdf" in h or "/article/download/" in h:
            full = urljoin(r.url, href.strip())
            if full not in seen:
                seen.add(full)
                out.append((full, title))
    if not out and re.search(r"/article/view/\d+/\d+", r.url):   # fallback
        out.append((r.url.replace("/article/view/", "/article/download/", 1), title))
    return out


def maybe_delete(path):
    """DELETE_PDF_AFTER = True ho to padhne ke baad PDF hata do."""
    if DELETE_PDF_AFTER and os.path.abspath(path).startswith(os.path.abspath(PDF_FOLDER)):
        try:
            os.remove(path)
        except OSError:
            pass


def download(session, url):
    name = os.path.basename(unquote(urlparse(url).path)) or "file.pdf"
    ojs = re.search(r"/article/download/(\d+)/(\d+)", url)
    if ojs:
        name = f"article_{ojs.group(1)}_{ojs.group(2)}.pdf"
    name = re.sub(r"[^\w.\-]+", "_", name)[:80]
    if not name.lower().endswith(".pdf"):
        name += ".pdf"
    path = os.path.join(PDF_FOLDER, f"{hashlib.md5(url.encode()).hexdigest()[:6]}_{name}")
    if os.path.exists(path) and os.path.getsize(path) > 0:
        return path                                   # pehle se download hai
    with session.get(url, timeout=TIMEOUT, stream=True) as r:
        r.raise_for_status()
        first = True
        with open(path + ".part", "wb") as f:
            for chunk in r.iter_content(65536):
                if first:
                    if b"%PDF" not in chunk[:1024]:
                        raise ValueError("Ye valid PDF nahi hai (site ne HTML diya - block ya galat link)")
                    first = False
                f.write(chunk)
        if first:
            raise ValueError("Khaali response aaya")
    os.replace(path + ".part", path)
    return path


# ---------- Excel ----------
def save_excel(results):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    results = sorted(results, key=lambda r: (r["idx"], r["pdf_url"]))
    emails_rows, article_rows, sno = [], [], 0
    for r in results:
        for x in r["rows"]:
            sno += 1
            emails_rows.append({
                "S.No": sno, "Article No.": r["idx"] + 1, "Article Title": r["title"],
                "Email": x["email"], "Page": x["pages"],
                "Check?": "Yes" if x["conf"] != "High" else "",
                "Note": x["method"] if x["conf"] != "High" else "",
                "PDF Link": r["pdf_url"],
            })
        if r["status"] == "OK" and not r["rows"]:
            status = "PDF padhi, par koi email nahi mila"
        elif r["status"] == "OK":
            status = "OK"
        else:
            status = f'{r["status"]}: {r["note"]}'
        article_rows.append({
            "Article No.": r["idx"] + 1, "Article Title": r["title"],
            "Total Emails": len(r["rows"]),
            "Emails": ", ".join(x["email"] for x in r["rows"]),
            "Status": status, "PDF Link": r["pdf_url"] or r["link"],
        })

    sheets = {
        "Emails": pd.DataFrame(emails_rows, columns=[
            "S.No", "Article No.", "Article Title", "Email", "Page", "Check?", "Note", "PDF Link"]),
        "Article_Wise": pd.DataFrame(article_rows, columns=[
            "Article No.", "Article Title", "Total Emails", "Emails", "Status", "PDF Link"]),
    }
    # Agar Excel file khuli hui hai (PermissionError), to naye naam se save karo
    path = OUTPUT_FILE
    try:
        open(path, "ab").close()
    except PermissionError:
        base, ext = os.path.splitext(OUTPUT_FILE)
        path = f"{base}_{time.strftime('%H%M%S')}{ext}"
        print(f"\n'{OUTPUT_FILE}' Excel me khuli hai, isliye naye naam se save kar raha hu: {path}")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            for c in df.columns:
                df[c] = df[c].map(clean)
            df.to_excel(xw, sheet_name=name, index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E78")
                c.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(ws.columns, 1):
                w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 60)
    return len(emails_rows), len({r["Email"] for r in emails_rows}), path


# ---------- Main ----------
def main():
    links = list(OrderedDict.fromkeys(l.strip() for l in LINKS if l and l.strip()))
    if not links:
        sys.exit("LINKS list khaali hai. Upar LINKS me apne links paste karo.")
    idx = {l: i for i, l in enumerate(links)}
    os.makedirs(PDF_FOLDER, exist_ok=True)
    session = make_session()
    results = []
    print(f"Total links: {len(links)}")

    try:
        jobs = []   # (link, pdf_url, title)
        with ThreadPoolExecutor(WORKERS) as ex:
            futs = {ex.submit(discover, session, l): l for l in links}
            for f in tqdm(as_completed(futs), total=len(futs), desc="Links check ho rahe hain"):
                link = futs[f]
                try:
                    found = f.result()
                    if not found:
                        raise ValueError("Is link par PDF nahi mila")
                    jobs += [(link, u, t) for u, t in found]
                except Exception as e:
                    results.append({"idx": idx[link], "link": link, "pdf_url": "", "title": "",
                                    "file": "", "status": "Link failed", "note": str(e),
                                    "total_pages": 0, "rows": []})

        with ThreadPoolExecutor(WORKERS) as ex:
            futs = {ex.submit(download, session, u): (l, u, t) for l, u, t in jobs}
            for f in tqdm(as_completed(futs), total=len(futs), desc="Download + emails nikal rahe hain"):
                link, url, title = futs[f]
                base = {"idx": idx[link], "link": link, "pdf_url": url, "title": title}
                try:
                    path = f.result()
                except Exception as e:
                    results.append({**base, "file": "", "status": "Download failed",
                                    "note": str(e), "total_pages": 0, "rows": []})
                    continue
                res = process_pdf(path)
                maybe_delete(path)
                results.append({**base, "file": os.path.basename(path), "status": res["status"],
                                "note": res["note"], "total_pages": res["total_pages"],
                                "rows": res["rows"]})
    except KeyboardInterrupt:
        print("\nRoka gaya - ab tak ka result save kar raha hu...")
    finally:
        if results:
            total, unique, out_path = save_excel(results)
            ok = sum(1 for r in results if r["status"] in ("OK", "No text"))
            bad = len(results) - ok
            print(f"\nKHATAM. {ok} PDFs padhi gayi, {bad} fail hui.")
            print(f"Total {total} email rows ({unique} alag-alag email).")
            print("Excel file:", os.path.abspath(out_path))


if __name__ == "__main__":
    main()